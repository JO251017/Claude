import glob
import os
import re
import sqlite3
import time
import zipfile
from datetime import date

import pytest
from openpyxl import load_workbook

from conftest import AS_OF, FAIL_KW, ROOT
from src import cli
from src.approval import ApprovalStructureError, import_approvals
from src.common.config import load_config
from src.common.db import init_db
from src.common.notify import Notifier
from src.discovery.trends import CachedTrendSource, CsvTrendSource, MockTrendSource, TrendFetchError


class Capture(Notifier):
    def __init__(self):
        self.items = []

    def notify(self, message, level="info", run_id=None):
        self.items.append((level, message))


def mock_source(env, extra_fail=None):
    """CSV 값을 그대로 쓰되 ratio를 식별 가능한 유일값으로 바꾼 모킹 소스 + 실패 주입."""
    csv_src = CsvTrendSource(env["files"]["trends"], date.fromisoformat(AS_OF))
    data, uniq = {}, []
    import csv
    kws = [r[0] for r in csv.reader(open(env["files"]["keywords"], encoding="utf-8")) if r and not r[0].startswith("#")][1:]
    for i, kw in enumerate(dict.fromkeys(kws)):
        try:
            series = csv_src.fetch(kw)
        except TrendFetchError:
            continue
        data[kw] = [None if v is None else round(7000.1234 + i * 13.0007 + j * 0.0173, 4) for j, v in enumerate(series)]
        uniq += [v for v in data[kw] if v is not None]
    failures = {FAIL_KW: TrendFetchError("server", "주입된 실패")}
    failures.update(extra_fail or {})
    return MockTrendSource(data, failures), uniq


def run_sample(env, source=None, notifier=None, keywords=None, run_id="RUN-E2E"):
    cfg = load_config(env["cfg_path"])
    notifier = notifier or Capture()
    source = source or mock_source(env)[0]
    res = cli.execute_run(cfg, keywords or env["files"]["keywords"], env["files"]["competitors"],
                          env["files"]["costs"], CachedTrendSource(source), date.fromisoformat(AS_OF), notifier,
                          run_id=run_id)
    return cfg, res, notifier


def test_e2e_time_and_outputs(sample_env):
    t0 = time.perf_counter()
    cfg, res, notifier = run_sample(sample_env)
    elapsed = time.perf_counter() - t0
    print(f"\n[기준① 샘플 20개 키워드 end-to-end 소요 {elapsed:.2f}초]")
    assert elapsed < 300
    files = glob.glob(str(sample_env["out"] / "*.xlsx"))
    assert len(files) == 1
    wb = load_workbook(files[0])
    assert wb.sheetnames == ["요약", "후보", "탈락·실패", "메타"] and wb["메타"].sheet_state == "hidden"
    meta = {r[0].value: r[1].value for r in wb["메타"].iter_rows(min_col=1, max_col=2)}
    assert meta["run_id"] == "RUN-E2E" and meta["source_name"] == "mock" and len(meta["config_hash"]) == 64
    f = res["facts"]
    assert f["pass"] >= 1 and f["hold"] >= 1 and f["reject"] >= 1 and f["input"] == 21
    assert notifier.items[-1] == ("info", f["conclusion"])


def test_sample_covers_required_cases(sample_env):
    cfg, res, _ = run_sample(sample_env)
    conn = sqlite3.connect(cfg.db_path); conn.row_factory = sqlite3.Row
    rows = {r["keyword"] + "|" + r["supplier"]: r for r in conn.execute("SELECT * FROM results")}
    reasons = " ".join(r["reasons"] for r in rows.values())
    for needle in ("보수 시나리오 순손실", "원가 미입력", "시세 미입력", "성장률 조회 실패", "최신성 초과 - 참고용"):
        assert needle in reasons, needle
    assert any(r["verdict"] == "탈락" for r in rows.values())
    # 기준③: 조회 실패 키워드가 결과에 포함되고 성장률 칸에 '조회 실패'
    fail = [r for r in rows.values() if r["keyword"] == FAIL_KW]
    assert fail and fail[0]["growth_status"] == "조회 실패"
    wb = load_workbook(res["xlsx"])
    grow = [r[4].value for r in wb["후보"].iter_rows(min_row=2) if r[0].value == FAIL_KW]
    assert grow == ["조회 실패"]
    print("[기준③ 조회 실패 키워드가 후보 시트 성장률 칸에 '조회 실패'로 표시됨]")


def test_sort_order(sample_env):
    _, res, _ = run_sample(sample_env)
    ws = load_workbook(res["xlsx"])["후보"]
    verdicts = [r[11].value for r in ws.iter_rows(min_row=2)]
    assert verdicts == sorted(verdicts, key=lambda v: 0 if v == "통과" else 1)
    passed = [r for r in ws.iter_rows(min_row=2) if r[11].value == "통과"]
    g = [r[4].value for r in passed]
    assert g == sorted(g, reverse=True)


def test_no_duplicate_trend_calls_and_no_timeseries_saved(sample_env, monkeypatch):
    monkeypatch.setenv("NAVER_HUB_CLIENT_SECRET", "DUMMYSECRET-ZZZ-9981")
    src, uniq = mock_source(sample_env)
    kw_file = sample_env["tmp"] / "kw.csv"
    kws = open(sample_env["files"]["keywords"], encoding="utf-8").read().splitlines()
    kw_file.write_text("\n".join(kws + [kws[3], kws[3]]) + "\n", encoding="utf-8")   # 같은 키워드를 중복 입력
    cfg, res, notifier = run_sample(sample_env, source=src, keywords=str(kw_file))
    assert len(src.calls) == len(set(src.calls)) == 20
    print(f"[검사] 같은 실행 중복 호출 {len(src.calls) - len(set(src.calls))}건")

    blobs = [open(cfg.db_path, "rb").read()]
    for p in glob.glob(str(sample_env["out"] / "**" / "*"), recursive=True):
        if os.path.isfile(p) and p.endswith(".xlsx"):
            with zipfile.ZipFile(p) as z:
                blobs += [z.read(n) for n in z.namelist()]
        elif os.path.isfile(p) and not p.endswith(".db"):
            blobs.append(open(p, "rb").read())
    blob = b"\n".join(blobs)
    assert uniq and "가상 접이식 정리함".encode() in blob     # 양성 대조: 검색 대상에 엑셀 본문이 실제로 포함됨
    for v in uniq:
        assert str(v).encode() not in blob
    # 비밀값 0건
    assert b"DUMMYSECRET-ZZZ-9981" not in blob
    print(f"[검사] ratio 값 {len(uniq)}개 / 비밀 더미값: DB·엑셀·로그·알림 파일에서 0건")


def test_secret_not_leaked_on_hub_path(sample_env, monkeypatch, capsys):
    import yaml
    monkeypatch.setenv("NAVER_HUB_CLIENT_ID", "DUMMYID-AAA-1")
    monkeypatch.setenv("NAVER_HUB_CLIENT_SECRET", "DUMMYSECRET-BBB-2")
    cfg = yaml.safe_load(open(sample_env["cfg_path"], encoding="utf-8"))
    cfg["hub"]["api_terms_ack"] = True
    open(sample_env["cfg_path"], "w", encoding="utf-8").write(yaml.safe_dump(cfg, allow_unicode=True))
    code = cli.main(["--config", sample_env["cfg_path"], "run", "--keywords", sample_env["files"]["keywords"],
                     "--competitors", sample_env["files"]["competitors"], "--costs", sample_env["files"]["costs"],
                     "--source", "hub", "--as-of", AS_OF])
    assert code == 4                                   # HubTrendSource 스텁 -> NotImplementedError
    out = capsys.readouterr().out
    assert "미구현" in out and "DUMMY" not in out
    blob = b"".join(open(p, "rb").read() for p in glob.glob(str(sample_env["out"] / "**" / "*"), recursive=True)
                    if os.path.isfile(p))
    assert b"DUMMY" not in blob


def test_hub_requires_ack_and_env_names_only(sample_env, monkeypatch, capsys):
    base = ["--config", sample_env["cfg_path"], "run", "--keywords", sample_env["files"]["keywords"],
            "--competitors", sample_env["files"]["competitors"], "--costs", sample_env["files"]["costs"], "--source", "hub"]
    assert cli.main(base) == 2
    assert "hub.api_terms_ack" in capsys.readouterr().out
    import yaml
    cfg = yaml.safe_load(open(sample_env["cfg_path"], encoding="utf-8")); cfg["hub"]["api_terms_ack"] = True
    open(sample_env["cfg_path"], "w", encoding="utf-8").write(yaml.safe_dump(cfg, allow_unicode=True))
    monkeypatch.delenv("NAVER_HUB_CLIENT_ID", raising=False); monkeypatch.delenv("NAVER_HUB_CLIENT_SECRET", raising=False)
    assert cli.main(base) == 2
    assert "NAVER_HUB_CLIENT_ID" in capsys.readouterr().out


def test_cli_missing_config_key_stops_before_run(tmp_path, capsys):
    bad = tmp_path / "c.yaml"; bad.write_text("fee_sales_pct: 3\n", encoding="utf-8")
    assert cli.main(["--config", str(bad), "run", "--keywords", "a", "--competitors", "b", "--costs", "c"]) == 2
    assert "fee_payment_pct" in capsys.readouterr().out


def test_error_notifications(sample_env):
    from src.discovery.trends import MockTrendSource
    cfg = load_config(sample_env["cfg_path"])
    # 전체 조회 실패
    allfail = MockTrendSource({}, {})
    _, res, n = run_sample(sample_env, source=allfail, run_id="RUN-F1")
    assert any(l == "error" and "전체" in m for l, m in n.items)
    # 한도 초과
    src, _ = mock_source(sample_env, extra_fail={"가상 주방 선반": TrendFetchError("rate_limit")})
    _, res, n = run_sample(sample_env, source=src, run_id="RUN-F2")
    assert any(l == "error" and "한도" in m for l, m in n.items)


def test_failure_logged_with_run_id_but_not_values(sample_env):
    cfg, _, _ = run_sample(sample_env)
    log = open(os.path.join(cfg.log_dir, "run_RUN-E2E.log"), encoding="utf-8").read()
    assert "조회 실패" in log and FAIL_KW in log and all("[RUN-E2E]" in ln for ln in log.splitlines() if ln)


def test_summary_consistency_stops_on_mismatch(sample_env):
    from src.report_xlsx import SummaryConsistencyError, build_summary, verify_summary
    from src.pipeline import split_rows, run_pipeline
    cfg, res, _ = run_sample(sample_env)
    # 일부러 어긋난 facts를 만들어 검증이 중단시키는지 확인
    from src.common.csvio import read_csv
    out = run_pipeline(cfg, [r["keyword"] for r in read_csv(sample_env["files"]["keywords"])],
                       read_csv(sample_env["files"]["competitors"]), read_csv(sample_env["files"]["costs"]),
                       mock_source(sample_env)[0], date.fromisoformat(AS_OF), "t")
    facts = build_summary(out.rows, 20)
    facts["pass"] += 1
    cands, rejects = split_rows(out.rows)
    with pytest.raises(SummaryConsistencyError):
        verify_summary(facts, out.rows, cands, rejects)


# ---------------- 승인 가져오기 ----------------
def _edit(xlsx, dst, run_val="승인", memo="메모", items=None, tamper=None):
    wb = load_workbook(xlsx)
    ws = wb["요약"]
    for row in ws.iter_rows(min_col=1, max_col=1):
        if row[0].value == "실행 승인":
            ws.cell(row[0].row, 2).value = run_val
        if row[0].value == "승인 메모":
            ws.cell(row[0].row, 2).value = memo
    wc = wb["후보"]
    col = [c.column for c in wc[1] if c.value == "항목 승인(O/X)"][0]
    for i, v in (items or {}).items():
        wc.cell(i, col).value = v
    if tamper == "run_id":
        wb["메타"]["B1"] = "OTHER-RUN"
    if tamper == "label":
        for row in ws.iter_rows(min_col=1, max_col=1):
            if row[0].value == "실행 승인":
                row[0].value = "x"
    wb.save(dst)
    return dst


def test_approval_import_flow(sample_env, tmp_path):
    cfg, res, _ = run_sample(sample_env)
    conn = init_db(cfg.db_path)
    f1 = _edit(res["xlsx"], str(tmp_path / "a1.xlsx"), items={2: "O", 3: "o", 4: "X", 5: "?", 6: None})
    rep = import_approvals(conn, f1)
    assert rep.inserted == 4 and len(rep.row_errors) == 1            # run + O + o + X, '?'는 건너뜀
    n_before = conn.execute("SELECT COUNT(*) FROM approval_events").fetchone()[0]
    rep2 = import_approvals(conn, f1)                                  # 같은 파일 재가져오기
    assert rep2.inserted == 0 and rep2.duplicates == 4
    assert conn.execute("SELECT COUNT(*) FROM approval_events").fetchone()[0] == n_before
    appr = conn.execute("SELECT keyword, supplier FROM approved_candidates ORDER BY keyword").fetchall()
    assert len(appr) == 2
    print(f"[공통기준 승인 가져오기] approval_events {n_before}건 기록, approved_candidates {len(appr)}건")
    # 실행 승인을 반려로 바꾼 새 파일 -> 뷰 비어짐
    f2 = _edit(res["xlsx"], str(tmp_path / "a2.xlsx"), run_val="반려", items={2: "O"})
    import_approvals(conn, f2)
    assert conn.execute("SELECT COUNT(*) FROM approved_candidates").fetchone()[0] == 0


def test_approval_structure_errors(sample_env, tmp_path):
    cfg, res, _ = run_sample(sample_env)
    conn = init_db(cfg.db_path)
    with pytest.raises(ApprovalStructureError):
        import_approvals(conn, _edit(res["xlsx"], str(tmp_path / "b1.xlsx"), tamper="run_id"))
    with pytest.raises(ApprovalStructureError):
        import_approvals(conn, _edit(res["xlsx"], str(tmp_path / "b2.xlsx"), tamper="label"))
    assert conn.execute("SELECT COUNT(*) FROM approval_events").fetchone()[0] == 0


# ---------------- 소스 검사 ----------------
def test_no_hardcoded_rates_in_src():
    pat = re.compile(r"14\.633|0\.14633|\b0\.15\b|\b2000\b|3\.003|3\.63\b")
    hits = []
    for p in glob.glob(os.path.join(ROOT, "src", "**", "*.py"), recursive=True):
        for i, ln in enumerate(open(p, encoding="utf-8"), 1):
            if pat.search(ln):
                hits.append(f"{p}:{i}")
    assert hits == []


def test_cli_csv_run_and_approve_import(sample_env, capsys, tmp_path):
    base = ["--config", sample_env["cfg_path"]]
    f = sample_env["files"]
    assert cli.main(base + ["run", "--keywords", f["keywords"], "--competitors", f["competitors"], "--costs", f["costs"],
                            "--trends", f["trends"], "--source", "csv", "--as-of", AS_OF]) == 0
    out = capsys.readouterr().out
    assert "통과" in out and "[info]" in out
    xlsx = glob.glob(str(sample_env["out"] / "*.xlsx"))[0]
    edited = _edit(xlsx, str(tmp_path / "ok.xlsx"), items={2: "O"})
    assert cli.main(base + ["approve", "import", edited]) == 0
    assert "신규 2건" in capsys.readouterr().out
    assert cli.main(base + ["approve", "import", _edit(xlsx, str(tmp_path / "bad.xlsx"), tamper="label")]) == 3
