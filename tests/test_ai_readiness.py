import os

import pytest
import yaml
from openpyxl import load_workbook

from src.common.config import ConfigError, load_config
from src.pipeline import ai_missing_items
from test_e2e import run_sample


def test_ai_missing_items_rules():
    full = {"spec_info": "Y", "model_name": "M-1", "dispatch_days": "0", "differentiator": "무타공"}
    assert ai_missing_items(full) == []
    assert ai_missing_items({**full, "spec_info": "N"}) == ["스펙 정보"]
    assert ai_missing_items({**full, "dispatch_days": "abc"}) == ["출고 소요일"]
    assert ai_missing_items({**full, "dispatch_days": "-1"}) == ["출고 소요일"]
    assert ai_missing_items({}) == ["스펙 정보", "모델명·브랜드", "출고 소요일", "차별점"]


def _cands(xlsx):
    ws = load_workbook(xlsx)["후보"]
    hdr = [c.value for c in ws[1]]
    return hdr, {r[0]: dict(zip(hdr, r)) for r in ws.iter_rows(min_row=2, values_only=True)}


def test_column_shown_and_verdict_unchanged_by_default(sample_env):
    _, res, _ = run_sample(sample_env)
    hdr, rows = _cands(res["xlsx"])
    assert hdr[-2:] == ["구분 라벨", "항목 승인(O/X)"] and "AI 노출 준비" in hdr
    assert rows["가상 수건 걸이"]["AI 노출 준비"] == "보완: 스펙 정보, 모델명·브랜드, 차별점"
    assert rows["가상 수건 걸이"]["판정"] == "통과"                     # affects_verdict=false
    assert rows["가상 접이식 정리함"]["AI 노출 준비"] == "준비됨"
    assert rows["가상 차량 방향제"]["AI 노출 준비"].startswith("해당 없음")
    assert set(res["facts"]["anomalies"]["AI 노출 정보 부족"]) == {"가상 수건 걸이", "가상 제습 박스", "가상 대형 박스 테이프"}


def test_affects_verdict_true_holds(sample_env):
    cfg = yaml.safe_load(open(sample_env["cfg_path"], encoding="utf-8"))
    cfg["ai_readiness"]["affects_verdict"] = True
    open(sample_env["cfg_path"], "w", encoding="utf-8").write(yaml.safe_dump(cfg, allow_unicode=True))
    _, res, _ = run_sample(sample_env)
    _, rows = _cands(res["xlsx"])
    r = rows["가상 수건 걸이"]
    assert r["판정"] == "보류" and "AI 노출 정보 부족 (스펙 정보, 모델명·브랜드, 차별점)" in r["사유"]


def test_config_key_required(sample_env):
    cfg = yaml.safe_load(open(sample_env["cfg_path"], encoding="utf-8"))
    del cfg["ai_readiness"]
    open(sample_env["cfg_path"], "w", encoding="utf-8").write(yaml.safe_dump(cfg, allow_unicode=True))
    with pytest.raises(ConfigError) as e:
        load_config(sample_env["cfg_path"])
    assert e.value.keys == ["ai_readiness.affects_verdict"]


def test_service_exposes_ai_missing(sample_env):
    from src import service
    run_sample(sample_env)
    s = service.get_run_summary("RUN-E2E", sample_env["cfg_path"])
    by = {c["keyword"]: c for c in s["candidates"]}
    assert by["가상 제습 박스"]["ai_missing"] == ["출고 소요일"] and by["가상 차량 방향제"]["ai_missing"] is None
