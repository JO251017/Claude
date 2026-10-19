import base64
import glob
import hashlib
import hmac
import json
import os
import urllib.error
import urllib.parse

import pytest
import yaml
from openpyxl import load_workbook

from src import cli
from src.common.config import load_config
from src.common.http_retry import HttpError
from src.discovery.keyword_finder import (FetchError, SearchAdClient, ShopSearchClient, _qc, discover, sign)
from test_e2e import Capture

ROOT = os.path.join(os.path.dirname(__file__), "..")
CFG = load_config(os.path.join(ROOT, "config.yaml"))


class Resp:
    def __init__(self, payload):
        self.payload = payload

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False

    def read(self):
        return json.dumps(self.payload).encode()


def test_signature_matches_official_formula():
    # 공식 샘플: base64(HMAC-SHA256(secret, "{ts}.{method}.{uri}"))
    expected = base64.b64encode(hmac.new(b"SEC", b"1700000000000.GET./keywordstool", hashlib.sha256).digest()).decode()
    assert sign("1700000000000", "GET", "/keywordstool", "SEC") == expected


def test_searchad_request_and_both_response_shapes():
    seen = {}

    def opener(req, timeout):
        seen["url"], seen["method"] = req.full_url, req.get_method()
        seen["h"] = {k.lower(): v for k, v in req.header_items()}
        return Resp(seen.get("payload"))

    c = SearchAdClient(CFG, "KEY", "SEC", "123", urlopen=opener, clock=lambda: 1700000000.0)
    seen["payload"] = {"keywordList": [{"relKeyword": "a"}]}
    assert c.related("접이식 정리함") == [{"relKeyword": "a"}]
    u = urllib.parse.urlsplit(seen["url"])
    assert f"{u.scheme}://{u.netloc}{u.path}" == "https://api.searchad.naver.com/keywordstool" and seen["method"] == "GET"
    assert urllib.parse.parse_qs(u.query) == {"hintKeywords": ["접이식 정리함"], "showDetail": ["1"]}
    h = seen["h"]
    assert h["x-timestamp"] == "1700000000000" and h["x-api-key"] == "KEY" and h["x-customer"] == "123"
    assert h["x-signature"] == sign("1700000000000", "GET", "/keywordstool", "SEC")
    seen["payload"] = [{"relKeyword": "b"}]                       # 명세상 배열 형태도 읽음
    assert c.related("x") == [{"relKeyword": "b"}]
    seen["payload"] = {"nope": 1}
    with pytest.raises(FetchError):
        c.related("x")


def test_shop_count_request_limit_and_errors():
    seen = {}

    def opener(req, timeout):
        seen["url"] = req.full_url
        seen["h"] = {k.lower(): v for k, v in req.header_items()}
        return Resp({"total": 12345, "items": []})

    import dataclasses
    s = ShopSearchClient(dataclasses.replace(CFG, disc_shop_daily_limit=1), "ID", "SECRET", urlopen=opener)
    assert s.count("수건 걸이") == 12345
    u = urllib.parse.urlsplit(seen["url"])
    assert f"{u.scheme}://{u.netloc}{u.path}" == "https://openapi.naver.com/v1/search/shop.json"
    assert urllib.parse.parse_qs(u.query) == {"query": ["수건 걸이"], "display": ["1"]}
    assert seen["h"]["x-naver-client-id"] == "ID"
    with pytest.raises(FetchError) as e:
        s.count("again")
    assert e.value.kind == "rate_limit"

    def err(req, timeout):
        raise urllib.error.HTTPError(req.full_url, 403, "x", {}, None)
    with pytest.raises(HttpError):
        ShopSearchClient(CFG, "i", "s", urlopen=err).count("k")


def test_qc_low_volume_not_zero():
    assert _qc("<10") == (None, True) and _qc(1234) == (1234, False) and _qc("560") == (560, False)


class FakeAd:
    def __init__(self, data, fail=()):
        self.data, self.fail, self.calls = data, set(fail), []

    def related(self, seed):
        self.calls.append(seed)
        if seed in self.fail:
            raise HttpError(400)
        return self.data[seed]


class FakeShop:
    def __init__(self, totals, fail=()):
        self.totals, self.fail, self.calls = totals, set(fail), 0

    def count(self, kw):
        self.calls += 1
        if kw in self.fail:
            raise FetchError("malformed")
        return self.totals[kw]


def kw(name, pc, mo, comp="low"):
    return {"relKeyword": name, "monthlyPcQcCnt": pc, "monthlyMobileQcCnt": mo, "compIdx": comp}


AD = {
    "정리함": [kw("정리함", 1000, 9000, "high"), kw("접이식 정리함", 300, 2700), kw("틈새 정리함", "<10", 50)],
    "선반": [kw("주방 선반", 500, 4500, "mid"), kw("접이식 정리함", 999, 999)],
}
SHOP = {"정리함": 500000, "접이식 정리함": 6000, "주방 선반": 4000}


def test_discover_merge_sort_and_status():
    import dataclasses
    seeds = [{"seed": "정리함", "category": "50000001"}, {"seed": "선반", "category": "50000002"},
             {"seed": "없는씨앗", "category": ""}]
    ad = FakeAd({**AD, "없는씨앗": []}, fail={"없는씨앗"})
    res = discover(CFG, seeds, ad, FakeShop(SHOP), sleep=lambda s: None, clock=lambda: 0.0)
    by = {r.keyword: r for r in res.rows}
    assert by["접이식 정리함"].seeds == ["정리함", "선반"] and by["접이식 정리함"].monthly == 3000   # 처음 값 유지
    assert by["접이식 정리함"].category == "50000001"
    assert by["틈새 정리함"].status == "검색수 10 미만" and by["틈새 정리함"].product_count is None
    assert res.seed_failures == {"없는씨앗": "client"}
    ok = [r.keyword for r in res.rows if r.status == "확인"]
    assert ok == ["주방 선반", "접이식 정리함", "정리함"]          # 경쟁강도 0.8 < 2.0 < 50
    assert float(by["주방 선반"].competition) == pytest.approx(0.8)


def test_discover_min_search_cap_and_shop_failure():
    import dataclasses
    cfg = dataclasses.replace(CFG, disc_min_monthly_search=4000, disc_max_candidates=1)
    res = discover(cfg, [{"seed": "정리함"}, {"seed": "선반"}], FakeAd(AD), FakeShop(SHOP, fail={"정리함"}),
                   sleep=lambda s: None, clock=lambda: 0.0)
    by = {r.keyword: r for r in res.rows}
    assert by["접이식 정리함"].status == "제외(검색수 기준 미달)"      # 3000 < 4000
    assert by["정리함"].status == "상품 수 조회 실패"                  # 검색수 1위라 상한 1개 안에 들어감
    assert by["주방 선반"].status == "미조회(상한 초과)"


def _cfg_file(tmp_path, ack=False):
    cfg = yaml.safe_load(open(os.path.join(ROOT, "config.yaml"), encoding="utf-8"))
    cfg["paths"] = {"db_path": str(tmp_path / "o" / "t.db"), "output_dir": str(tmp_path / "o"),
                    "log_dir": str(tmp_path / "o" / "logs")}
    cfg["discovery"]["api_terms_ack"] = ack
    p = tmp_path / "c.yaml"
    p.write_text(yaml.safe_dump(cfg, allow_unicode=True), encoding="utf-8")
    seeds = tmp_path / "seeds.csv"
    seeds.write_text("# 테스트\nseed,category\n정리함,50000001\n선반,50000002\n정리함,50000001\n", encoding="utf-8")
    return str(p), str(seeds)


def test_execute_discover_outputs(tmp_path):
    cfg_path, seeds = _cfg_file(tmp_path)
    cfg = load_config(cfg_path)
    n = Capture()
    ad = FakeAd(AD)
    out = cli.execute_discover(cfg, seeds, n, "RUN-D1", ad=ad, shop=FakeShop(SHOP), sleep=lambda s: None)
    assert ad.calls == ["정리함", "선반"]                                # 중복 씨앗은 한 번만
    lines = [l for l in open(out["csv"], encoding="utf-8").read().splitlines() if not l.startswith("#")]
    assert lines == ["keyword,category", "주방 선반,50000002", "접이식 정리함,50000001", "정리함,50000001"]
    from src.common.csvio import read_csv
    assert [r["keyword"] for r in read_csv(out["csv"])][0] == "주방 선반"   # 1단계 입력으로 바로 읽힘
    wb = load_workbook(out["xlsx"])
    assert wb.sheetnames == ["요약", "발굴 결과"]
    concl = [r[1].value for r in wb["요약"].iter_rows() if r[0].value == "결론"][0]
    assert "주방 선반" in concl and n.items[-1] == ("info", concl)


def test_cli_discover_requires_ack_and_keys(tmp_path, monkeypatch, capsys):
    cfg_path, seeds = _cfg_file(tmp_path, ack=False)
    assert cli.main(["--config", cfg_path, "discover", "--seeds", seeds]) == 2
    assert "discovery.api_terms_ack" in capsys.readouterr().out
    cfg_path, seeds = _cfg_file(tmp_path, ack=True)
    for k in ("NAVER_AD_API_KEY", "NAVER_AD_SECRET_KEY", "NAVER_AD_CUSTOMER_ID", "NAVER_CLIENT_ID", "NAVER_CLIENT_SECRET"):
        monkeypatch.delenv(k, raising=False)
    monkeypatch.chdir(tmp_path)
    assert cli.main(["--config", cfg_path, "discover", "--seeds", seeds]) == 2
    out = capsys.readouterr().out
    assert "NAVER_AD_API_KEY" in out and "NAVER_CLIENT_ID" in out


def test_cli_discover_end_to_end_mocked_no_secret_leak(tmp_path, monkeypatch, capsys):
    import urllib.request
    cfg_path, seeds = _cfg_file(tmp_path, ack=True)
    secrets = {"NAVER_AD_API_KEY": "DUMMY-ADKEY-1", "NAVER_AD_SECRET_KEY": "DUMMY-ADSEC-2",
               "NAVER_AD_CUSTOMER_ID": "DUMMY-CUST-3", "NAVER_CLIENT_ID": "DUMMY-CID-4", "NAVER_CLIENT_SECRET": "DUMMY-CSEC-5"}
    for k, v in secrets.items():
        monkeypatch.setenv(k, v)
    monkeypatch.chdir(tmp_path)

    def fake(req, timeout):
        u = urllib.parse.urlsplit(req.full_url)
        q = urllib.parse.parse_qs(u.query)
        if u.path == "/keywordstool":
            return Resp({"keywordList": AD[q["hintKeywords"][0]]})
        return Resp({"total": SHOP[q["query"][0]]})

    monkeypatch.setattr(urllib.request, "urlopen", fake)
    monkeypatch.setattr("time.sleep", lambda s: None)
    assert cli.main(["--config", cfg_path, "discover", "--seeds", seeds]) == 0
    printed = capsys.readouterr().out
    blob = printed.encode()
    for p in glob.glob(str(tmp_path / "o" / "**" / "*"), recursive=True):
        if os.path.isfile(p):
            if p.endswith(".xlsx"):
                import zipfile
                with zipfile.ZipFile(p) as z:
                    blob += b"".join(z.read(x) for x in z.namelist())
            else:
                blob += open(p, "rb").read()
    assert "주방 선반".encode() in blob                                   # 양성 대조
    assert b"DUMMY" not in blob
