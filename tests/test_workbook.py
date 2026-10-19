import glob
import os
import sqlite3
from datetime import datetime

import pytest
import yaml
from openpyxl import load_workbook

from src import cli
from src import workbook as wbk
from test_e2e import Capture
from test_keyword_finder import AD, SHOP, FakeAd, FakeShop

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
AS_OF = "2026-07-15"


@pytest.fixture
def env(tmp_path, monkeypatch):
    cfg = yaml.safe_load(open(os.path.join(ROOT, "data", "sample", "config.sample.yaml"), encoding="utf-8"))
    cfg["paths"] = {"db_path": str(tmp_path / "o" / "t.db"), "output_dir": str(tmp_path / "o"),
                    "log_dir": str(tmp_path / "o" / "logs")}
    (tmp_path / "c.yaml").write_text(yaml.safe_dump(cfg, allow_unicode=True), encoding="utf-8")
    monkeypatch.chdir(ROOT)                      # data/sample 상대경로 사용
    wb_path = str(tmp_path / "관리.xlsx")
    return {"cfg": str(tmp_path / "c.yaml"), "wb": wb_path, "tmp": tmp_path,
            "base": ["--config", str(tmp_path / "c.yaml"), "--workbook", wb_path]}


def test_init_creates_input_sheets_and_refuses_overwrite(env, capsys):
    assert cli.main(env["base"] + ["init"]) == 0
    wb = load_workbook(env["wb"])
    assert wb.sheetnames == ["안내", "씨앗", "키워드", "경쟁시세", "도매원가", "클릭추이"]
    assert [c.value for c in wb["도매원가"][1]][:3] == ["키워드", "상품명", "도매 원가"]
    assert cli.main(env["base"] + ["init"]) == 2
    assert "--force" in capsys.readouterr().out


def test_read_sheet_converts_excel_types(env):
    cli.main(env["base"] + ["init"])
    wb = load_workbook(env["wb"])
    ws = wb["경쟁시세"]
    ws.append(["k", "s", 12000.0, 0, datetime(2026, 7, 12)])
    wt = wb["클릭추이"]
    wt.append(["k", datetime(2025, 4, 1), 33.5])
    wb.save(env["wb"])
    wb = load_workbook(env["wb"])
    rows = wbk.read_sheet(wb, "경쟁시세")
    assert rows[-1] == {"keyword": "k", "seller": "s", "price": "12000", "shipping": "0", "researched_at": "2026-07-12"}
    assert wbk.read_sheet(wb, "클릭추이")[-1]["month"] == "2025-04"          # 엑셀이 날짜로 바꿔도 월로 읽음


def test_run_in_workbook_matches_csv_run(env):
    assert cli.main(env["base"] + ["init", "--sample"]) == 0
    assert cli.main(env["base"] + ["run", "--as-of", AS_OF]) == 0
    wb = load_workbook(env["wb"])
    assert wb.sheetnames[:6] == ["안내", "씨앗", "키워드", "경쟁시세", "도매원가", "클릭추이"]
    assert {"요약", "후보", "탈락·실패", "메타"} <= set(wb.sheetnames) and wb["메타"].sheet_state == "hidden"
    assert wb["도매원가"].max_row >= 22                                       # 입력은 그대로 남음
    concl_wb = [r[1].value for r in wb["요약"].iter_rows() if r[0].value == "결론"][0]
    # 같은 샘플을 CSV 방식으로 돌린 결과와 같아야 한다
    s = os.path.join(ROOT, "data", "sample")
    assert cli.main(["--config", env["cfg"], "run", "--keywords", f"{s}/keywords.csv", "--competitors",
                     f"{s}/competitors.csv", "--costs", f"{s}/costs.csv", "--trends", f"{s}/trends.csv",
                     "--as-of", AS_OF]) == 0
    csv_x = sorted(glob.glob(str(env["tmp"] / "o" / "candidates_*.xlsx")), key=os.path.getmtime)[-1]
    concl_csv = [r[1].value for r in load_workbook(csv_x)["요약"].iter_rows() if r[0].value == "결론"][0]
    assert concl_wb == concl_csv


def test_approve_from_workbook_default(env, capsys):
    cli.main(env["base"] + ["init", "--sample"])
    cli.main(env["base"] + ["run", "--as-of", AS_OF])
    wb = load_workbook(env["wb"])
    ws = wb["요약"]
    for r in ws.iter_rows(min_col=1, max_col=1):
        if r[0].value == "실행 승인":
            ws.cell(r[0].row, 2).value = "승인"
    wc = wb["후보"]
    col = [c.column for c in wc[1] if c.value == "항목 승인(O/X)"][0]
    wc.cell(2, col).value = "O"
    wb.save(env["wb"])
    capsys.readouterr()
    assert cli.main(env["base"] + ["approve"]) == 0
    assert "신규 2건" in capsys.readouterr().out
    conn = sqlite3.connect(str(env["tmp"] / "o" / "t.db"))
    assert conn.execute("SELECT COUNT(*) FROM approved_candidates").fetchone()[0] == 1


def test_discover_into_workbook_keeps_marks_and_feeds_run(env):
    from src.common.config import load_config
    cli.main(env["base"] + ["init"])
    cfg = load_config(env["cfg"])
    seeds = [{"seed": "정리함", "category": "50000001"}, {"seed": "선반", "category": "50000002"}]
    cli.execute_discover(cfg, seeds, Capture(), "D1", ad=FakeAd(AD), shop=FakeShop(SHOP),
                         sleep=lambda s: None, workbook_path=env["wb"])
    wb = load_workbook(env["wb"])
    assert wb.sheetnames.index("발굴 결과") == wb.sheetnames.index("씨앗") + 1
    ws = wb["발굴 결과"]
    hdr = [c.value for c in ws[1]]
    ws.cell(2, hdr.index("키워드로 사용(O)") + 1).value = "O"               # 첫 줄(주방 선반) 선택
    wb["키워드"].append(["직접 넣은 키워드", "123"])
    wb.save(env["wb"])
    cli.execute_discover(cfg, seeds, Capture(), "D2", ad=FakeAd(AD), shop=FakeShop(SHOP),
                         sleep=lambda s: None, workbook_path=env["wb"])   # 다시 발굴해도 O 유지
    wb = load_workbook(env["wb"])
    assert wbk.selected_discovered(wb) == [{"keyword": "주방 선반", "category": "50000002"}]
    assert [r["keyword"] for r in wbk.keywords_for_run(wb)] == ["직접 넣은 키워드", "주방 선반"]


def test_locked_workbook_stops_before_recording(env, monkeypatch, capsys):
    cli.main(env["base"] + ["init", "--sample"])

    def locked(self, path):
        raise PermissionError("locked")
    monkeypatch.setattr("openpyxl.workbook.workbook.Workbook.save", locked)
    assert cli.main(env["base"] + ["run", "--as-of", AS_OF]) == 5
    assert "엑셀을 닫고" in capsys.readouterr().out
    conn = sqlite3.connect(str(env["tmp"] / "o" / "t.db"))
    assert conn.execute("SELECT status FROM runs").fetchall() == [("failed",)]
    assert conn.execute("SELECT COUNT(*) FROM results").fetchone()[0] == 0
