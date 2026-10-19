"""엑셀 출력(openpyxl). 결론·근거는 코드 템플릿으로 생성하며 AI를 쓰지 않는다."""
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

NO_MARKET_LABEL = "시세 미입력"
from src.pipeline import INPUT_ERROR, ResultRow, split_rows
from src.workbook import replace_sheet
from src.xlsx_style import style_candidates, style_rejects, style_summary
from src.sourcing.profit import LABEL, REASON_LABELS, TAX_NOTE, round_half_up

SHEET_SUMMARY, SHEET_CANDS, SHEET_REJECTS, SHEET_META = "요약", "후보", "탈락·실패", "메타"
LBL_RUN_ID, LBL_RUN_APPROVAL, LBL_RUN_MEMO = "run_id", "실행 승인", "승인 메모"
CAND_HEADERS = ["키워드", "도매처", "시장 기준가", "판매처 수", "클릭 증가율", "원가", "도매 배송비",
                "순이익(세전: 부가세·소득세 미반영)", "이익률", "보수 순이익", "광고비 손익분기율", "판정", "사유", "경고",
                "AI 노출 준비", "구분 라벨", "항목 승인(O/X)"]
REJECT_HEADERS = ["키워드", "도매처", "구분", "부족액(원)", "사유", "실제값 대 기준값"]
HDR_FILL = PatternFill("solid", fgColor="DDE6F0")
WON, PCT = "#,##0", "0.0%"


class SummaryConsistencyError(Exception):
    pass


def _distinct(rows, pred):
    return sorted({r.keyword for r in rows if pred(r)})


def build_summary(rows: list[ResultRow], total_keywords: int) -> dict:
    cands, rejects = split_rows(rows)
    cnt = Counter(r.verdict for r in rows)
    reason_cnt: Counter = Counter()
    for r in rows:
        if r.verdict in ("보류", "탈락", INPUT_ERROR):
            for code in {x.code for x in r.reasons}:
                reason_cnt[(REASON_LABELS[code], r.verdict)] += 1
    passed = [r for r in cands if r.verdict == "통과"]
    top = passed[0] if passed else None

    facts = {
        "input": len(rows), "pass": cnt["통과"], "hold": cnt["보류"], "reject": cnt["탈락"],
        "input_error": cnt[INPUT_ERROR], "reason_dist": reason_cnt, "total_keywords": total_keywords,
        "top": top,
        "anomalies": {
            "조회 실패": _distinct(rows, lambda r: r.growth is not None and r.growth.status == "조회 실패"),
            "시세 미입력": _distinct(rows, lambda r: any(x.code == "NO_MARKET" for x in r.reasons)),
            "원가 미입력": _distinct(rows, lambda r: any(x.code == "NO_COST" for x in r.reasons)),
            "최신성 초과": _distinct(rows, lambda r: any(x.code == "STALE" for x in r.reasons)),
            "표본 부족": _distinct(rows, lambda r: any(w.startswith("표본 부족") for w in r.warnings)),
            "AI 노출 정보 부족": _distinct(rows, lambda r: r.verdict in ("통과", "보류") and bool(r.ai_missing)),
        },
    }
    if top is not None:
        facts["conclusion"] = (f"통과 {facts['pass']}건 · 보류 {facts['hold']}건 · 탈락 {facts['reject']}건"
                               f" (입력 오류 {facts['input_error']}건). 최우선 후보: {top.keyword} / {top.supplier}, "
                               f"순이익 {top.profit.profit:,}원(세전)")
    else:
        facts["conclusion"] = (f"통과 후보 없음 - 보류 {facts['hold']}건 · 탈락 {facts['reject']}건"
                               f" (입력 오류 {facts['input_error']}건)")

    ev = []
    ev.append((f"통과 {facts['pass']}건은 후보 시트에 클릭 증가율 내림차순 → 순이익 내림차순으로 정렬", SHEET_CANDS))
    if reason_cnt:
        (label, verdict), n = reason_cnt.most_common(1)[0]
        ev.append((f"탈락·보류·입력 오류 사유 최다: {label}({verdict}) {n}건", SHEET_REJECTS + "/" + SHEET_CANDS))
    n_anom = sum(len(v) for v in facts["anomalies"].values())
    ev.append((f"이상 징후 {n_anom}건(조회 실패 {len(facts['anomalies']['조회 실패'])}, 시세·원가 미입력 "
               f"{len(facts['anomalies']['시세 미입력']) + len(facts['anomalies']['원가 미입력'])}, 최신성 초과 "
               f"{len(facts['anomalies']['최신성 초과'])}, 표본 부족 {len(facts['anomalies']['표본 부족'])}, "
               f"AI 노출 정보 부족 {len(facts['anomalies']['AI 노출 정보 부족'])})", SHEET_SUMMARY))
    facts["evidence"] = ev
    facts["n_cands"], facts["n_rejects"] = len(cands), len(rejects)
    verify_summary(facts, rows, cands, rejects)
    return facts


def verify_summary(facts, rows, cands, rejects) -> None:
    """근거 수치가 시트에 실제로 들어갈 행과 맞는지 독립 재집계로 확인. 어긋나면 생성을 중단한다."""
    p = sum(1 for r in cands if r.verdict == "통과")
    h = sum(1 for r in cands if r.verdict == "보류")
    j = sum(1 for r in rejects if r.verdict == "탈락")
    e = sum(1 for r in rejects if r.verdict == INPUT_ERROR)
    if (p, h, j, e) != (facts["pass"], facts["hold"], facts["reject"], facts["input_error"]):
        raise SummaryConsistencyError(f"요약 수치 불일치: 시트 {(p, h, j, e)} vs 요약 "
                                      f"{(facts['pass'], facts['hold'], facts['reject'], facts['input_error'])}")
    if p + h + j + e != facts["input"] or facts["input"] != len(rows):
        raise SummaryConsistencyError("요약 수치 불일치: 입력 건수 != 통과+보류+탈락+입력 오류")
    first_pass = next((r for r in cands if r.verdict == "통과"), None)
    if first_pass is not facts["top"]:
        raise SummaryConsistencyError("요약 수치 불일치: 최우선 후보가 후보 시트 1순위와 다름")
    if facts["top"] is not None and (facts["top"].keyword not in facts["conclusion"]):
        raise SummaryConsistencyError("요약 수치 불일치: 결론에 최우선 후보 없음")


def write_workbook(path: str, rows: list[ResultRow], facts: dict, meta: dict) -> None:
    """결과만 담은 독립 엑셀(실행 기록 사본)."""
    wb = Workbook()
    fill_result_sheets(wb, rows, facts, meta)
    del wb[wb.sheetnames[0]]
    wb.save(path)


def fill_result_sheets(wb, rows: list[ResultRow], facts: dict, meta: dict) -> None:
    """요약·후보·탈락·실패·메타 시트를 wb에 새로 쓴다(같은 이름 시트는 교체). 관리 엑셀에도 그대로 쓴다."""
    cands, rejects = split_rows(rows)

    # ---- 요약 ----
    ws = replace_sheet(wb, SHEET_SUMMARY)
    bold = Font(bold=True)
    r = 1
    ws.cell(r, 1, "순이익 후보표 - 요약").font = Font(bold=True, size=14); r += 1
    ws.cell(r, 1, LBL_RUN_ID); ws.cell(r, 2, meta["run_id"]); r += 1
    ws.cell(r, 1, "기준"); ws.cell(r, 2, f"순이익은 {TAX_NOTE}"); r += 2
    ws.cell(r, 1, "결론").font = bold; ws.cell(r, 2, facts["conclusion"]); r += 1
    for i, (text, sheet) in enumerate(facts["evidence"], 1):
        ws.cell(r, 1, f"근거 {i}").font = bold; ws.cell(r, 2, text); ws.cell(r, 3, f"확인: {sheet} 시트"); r += 1
    r += 1
    ws.cell(r, 1, "처리 수치").font = bold; r += 1
    for lab, key in (("입력", "input"), ("통과", "pass"), ("보류", "hold"), ("탈락", "reject"), ("입력 오류", "input_error")):
        ws.cell(r, 1, lab); ws.cell(r, 2, facts[key]); r += 1
    r += 1
    ws.cell(r, 1, "탈락·보류 사유 분포").font = bold; r += 1
    for (label, verdict), n in sorted(facts["reason_dist"].items(), key=lambda kv: (-kv[1], kv[0])):
        ws.cell(r, 1, label); ws.cell(r, 2, n); ws.cell(r, 3, verdict); r += 1
    if not facts["reason_dist"]:
        ws.cell(r, 1, "해당 없음"); r += 1
    r += 1
    ws.cell(r, 1, "이상 징후").font = bold; r += 1
    for lab, kws in facts["anomalies"].items():
        ws.cell(r, 1, lab); ws.cell(r, 2, len(kws)); ws.cell(r, 3, ", ".join(kws[:8]) + (" ..." if len(kws) > 8 else "")); r += 1
    r += 1
    ws.cell(r, 1, "결정 요청").font = bold
    ws.cell(r, 2, f"통과·보류 후보 {facts['n_cands']}건의 실행 여부를 아래 칸에 입력하고, 항목별 O/X는 후보 시트에 입력하세요.")
    r += 2
    ws.cell(r, 1, "승인 입력").font = bold; r += 1
    ws.cell(r, 1, LBL_RUN_APPROVAL).font = bold
    dv = DataValidation(type="list", formula1='"승인,반려,보류"', allow_blank=True)
    ws.add_data_validation(dv); dv.add(ws.cell(r, 2)); r += 1
    ws.cell(r, 1, LBL_RUN_MEMO).font = bold
    ws.column_dimensions["A"].width = 24; ws.column_dimensions["B"].width = 80; ws.column_dimensions["C"].width = 40

    # ---- 후보 ----
    wc = replace_sheet(wb, SHEET_CANDS)
    wc.append(CAND_HEADERS)
    for row in cands:
        m, pr = row.market, row.profit
        wc.append([
            row.keyword, row.supplier,
            round_half_up(m.price) if m else NO_MARKET_LABEL, m.seller_count if m else None,
            row.growth_display,
            round_half_up(row.cost) if row.cost is not None else None,
            round_half_up(row.ship) if row.ship is not None else None,
            pr.profit if pr else row.profit_text,
            pr.margin if pr else None, pr.stress_profit if pr else None, pr.a_star if pr else None,
            row.verdict, "; ".join(x.text for x in row.reasons), "; ".join(row.warnings),
            ai_readiness_text(row.ai_missing), LABEL, None,
        ])
    _format_candidates(wc, len(cands))

    # ---- 탈락·실패 ----
    wr = replace_sheet(wb, SHEET_REJECTS)
    wr.append(REJECT_HEADERS)
    for row in rejects:
        pr = row.profit
        detail = ""
        if pr and not pr.input_error:
            parts = []
            if any(x.code == "BELOW_TMIN" for x in row.reasons):
                parts.append(f"T_sug {pr.t_sug:,} vs T_min {pr.t_min:,}")
            if any(x.code == "P_NONPOS" for x in row.reasons):
                parts.append(f"P {pr.p:,} vs 기준 0 초과")
            detail = "; ".join(parts)
        elif row.verdict == "탈락":
            detail = "성장률 하한 미달"
        wr.append([row.keyword, row.supplier, row.verdict, pr.shortfall if pr else None,
                   "; ".join(x.text for x in row.reasons), detail])
    for col, w in zip("ABCDEF", (16, 14, 10, 12, 60, 40)):
        wr.column_dimensions[col].width = w
    for c in wr[1]: c.font = bold; c.fill = HDR_FILL

    # ---- 메타(숨김) ----
    wm = replace_sheet(wb, SHEET_META)
    for k in ("run_id", "config_hash", "source_name", "executed_at", "as_of"):
        wm.append([k, meta[k]])
    wm.sheet_state = "hidden"

    style_summary(ws)
    style_candidates(wc, {h: i for i, h in enumerate(CAND_HEADERS, 1)})
    style_rejects(wr)




def ai_readiness_text(missing) -> str:
    if missing is None:
        return "해당 없음(원가 미입력)"
    return "준비됨" if not missing else "보완: " + ", ".join(missing)


def _format_candidates(ws, n):
    widths = [16, 14, 12, 10, 12, 10, 12, 22, 9, 12, 14, 8, 50, 40, 24, 36, 14]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
        ws.cell(1, i).font = Font(bold=True); ws.cell(1, i).fill = HDR_FILL
        ws.cell(1, i).alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C2"
    dv = DataValidation(type="list", formula1='"O,X"', allow_blank=True)
    ws.add_data_validation(dv)
    for i in range(2, n + 2):
        for col, fmt in ((3, WON), (6, WON), (7, WON), (8, WON), (10, WON), (9, PCT), (11, PCT), (5, PCT)):
            ws.cell(i, col).number_format = fmt
        dv.add(ws.cell(i, len(CAND_HEADERS)))
