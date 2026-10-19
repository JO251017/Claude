"""모듈 0 결과: 엑셀(요약, 발굴 결과) + 1단계 입력용 keywords.csv."""
import csv
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Alignment

from src.discovery.keyword_finder import DiscoveryResult
from src.workbook import DISCOVERY, USE_HEADER, replace_sheet, selected_discovered
from src.xlsx_style import BOX, INPUT, MUTED, NAVY, VERDICT, _fill, _font, _header

HEADERS = ["키워드", "씨앗 키워드", "분야 코드", "월간 검색수(PC+모바일)", "PC", "모바일", "상품 수",
           "경쟁강도(상품 수÷검색수)", "광고 경쟁정도", "상태", USE_HEADER]


def conclusion(res: DiscoveryResult) -> str:
    ok = [r for r in res.rows if r.status == "확인"]
    if not ok:
        return f"연관 키워드 {len(res.rows)}개, 상품 수까지 확인된 키워드 없음"
    top = ok[0]
    return (f"연관 키워드 {len(res.rows)}개 중 {len(ok)}개 확인. 경쟁강도 가장 낮은 키워드: {top.keyword} "
            f"(월간 검색 {top.monthly:,}, 상품 {top.product_count:,})")


def write_keywords_csv(path: str, res: DiscoveryResult) -> int:
    rows = [r for r in res.rows if r.status == "확인"]
    with open(path, "w", encoding="utf-8", newline="") as f:
        f.write("# 키워드 발굴 결과(경쟁강도 낮은 순). 필요 없는 줄을 지우고 data/input/keywords.csv로 복사하세요.\n")
        w = csv.writer(f)
        w.writerow(["keyword", "category"])
        for r in rows:
            w.writerow([r.keyword, r.category])
    return len(rows)


def write_discovery_xlsx(path: str, res: DiscoveryResult, meta: dict, csv_name: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "요약"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 90
    lines = [("키워드 발굴 - 요약", ""), ("run_id", meta["run_id"]), ("실행 시각", meta["executed_at"]), ("", ""),
             ("결론", conclusion(res)), ("", "")]
    status_cnt = Counter(r.status for r in res.rows)
    lines += [("처리 수치", "")] + [(k, v) for k, v in [
        ("씨앗 키워드", meta["n_seeds"]), ("연관 키워드", len(res.rows)), ("상품 수 확인", status_cnt["확인"]),
        ("상품 수 조회 실패", status_cnt["상품 수 조회 실패"]), ("검색수 10 미만", status_cnt["검색수 10 미만"]),
        ("미조회(상한 초과)", status_cnt["미조회(상한 초과)"]), ("제외(검색수 기준 미달)", status_cnt["제외(검색수 기준 미달)"]),
        ("쇼핑 검색 호출 수", res.shop_calls)]]
    lines += [("", ""), ("조회 실패한 씨앗", ", ".join(f"{s}({k})" for s, k in res.seed_failures.items()) or "없음"),
              ("", ""), ("다음 단계", f"{csv_name}을 열어 팔 키워드만 남기고 data/input/keywords.csv로 복사한 뒤 "
                                   "python -m src.cli run 을 실행하세요."),
              ("읽는 법", "경쟁강도 = 상품 수 ÷ 월간 검색수. 낮을수록 찾는 사람에 비해 파는 상품이 적습니다. "
                       "광고 경쟁정도는 검색광고 기준(low/mid/high)이며 판정에는 쓰지 않습니다.")]
    for a, b in lines:
        ws.append([a, b])
    ws["A1"].font = _font(size=16, bold=True, color=NAVY)
    for row in ws.iter_rows(min_row=2):
        a, b = row[0], row[1]
        a.alignment = Alignment(vertical="center")
        b.alignment = Alignment(wrap_text=True, vertical="center")
        if a.value in ("처리 수치",):
            a.fill = b.fill = _fill(NAVY)
            a.font = _font(bold=True, color="FFFFFF")
        elif a.value in ("결론", "다음 단계"):
            a.fill = b.fill = _fill("E8F0FA")
            a.font = _font(bold=True, color=NAVY)
            b.font = _font(bold=a.value == "결론", size=12 if a.value == "결론" else 10, color=NAVY)
            ws.row_dimensions[a.row].height = 36
        elif a.value:
            a.font = _font(bold=True)
            b.font = _font(color=MUTED if a.value in ("run_id", "실행 시각", "읽는 법") else "222B36")

    wr = fill_discovery_results(wb, res)
    for sheet, orient in ((ws, "portrait"), (wr, "landscape")):
        sheet.page_setup.orientation = orient
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(path)


def fill_discovery_results(wb, res: DiscoveryResult, index: int | None = None):
    """'발굴 결과' 시트를 새로 쓴다. 이미 O 표시한 키워드는 그대로 O로 남긴다."""
    marked = {r["keyword"] for r in selected_discovered(wb)}
    wr = replace_sheet(wb, DISCOVERY, index)
    wr.append(HEADERS)
    for r in res.rows:
        low = lambda v: "<10" if v is None and r.low_volume else v  # noqa: E731
        wr.append([r.keyword, ", ".join(r.seeds), r.category, r.monthly, low(r.pc), low(r.mobile), r.product_count,
                   float(r.competition) if r.competition is not None else None, r.comp_idx, r.status,
                   "O" if r.keyword in marked else None])
    _header(wr, 36)
    wr.sheet_view.showGridLines = False
    wr.freeze_panes = "B2"
    wr.auto_filter.ref = wr.dimensions
    for col, w in zip("ABCDEFGHIJK", (24, 22, 12, 14, 10, 10, 12, 14, 12, 20, 14)):
        wr.column_dimensions[col].width = w
    for row in wr.iter_rows(min_row=2):
        for c in row:
            c.font = _font()
            c.border = BOX
        row[0].font = _font(bold=True)
        for c in row[3:7]:
            c.number_format = "#,##0"
        row[7].number_format = "0.00"
        st = row[9]
        bg, fg = VERDICT["통과"] if st.value == "확인" else (VERDICT["보류"] if st.value != "제외(검색수 기준 미달)" else VERDICT["입력 오류"])
        st.fill, st.font = _fill(bg), _font(bold=True, color=fg)
        use = row[10]
        use.fill, use.font = _fill(INPUT), _font(bold=True)
        use.alignment = Alignment(horizontal="center")
    from openpyxl.worksheet.datavalidation import DataValidation
    dv = DataValidation(type="list", formula1='"O"', allow_blank=True)
    wr.add_data_validation(dv)
    dv.add(f"K2:K{max(wr.max_row, 2)}")
    wr.page_setup.orientation = "landscape"
    wr.page_setup.fitToWidth = 1
    wr.page_setup.fitToHeight = 0
    wr.sheet_properties.pageSetUpPr.fitToPage = True
    return wr
