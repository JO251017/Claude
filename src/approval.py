"""승인 가져오기: 요약 시트의 실행 승인 칸과 후보 시트의 O/X를 approval_events에 기록."""
import hashlib
from dataclasses import dataclass, field
from datetime import datetime

from openpyxl import load_workbook

from src.report_xlsx import (CAND_HEADERS, LBL_RUN_APPROVAL, LBL_RUN_ID, LBL_RUN_MEMO, SHEET_CANDS, SHEET_META,
                             SHEET_SUMMARY)

RUN_DECISIONS = {"승인", "반려", "보류"}
ITEM_MAP = {"O": "승인", "X": "반려"}
ITEM_HEADER = CAND_HEADERS[-1]


class ApprovalStructureError(Exception):
    pass


@dataclass
class ImportReport:
    run_id: str = ""
    inserted: int = 0
    duplicates: int = 0
    row_errors: list = field(default_factory=list)


def _find_label(ws, label):
    for row in ws.iter_rows(min_col=1, max_col=1):
        c = row[0]
        if c.value is not None and str(c.value).strip() == label:
            return c.row
    return None


def import_approvals(conn, xlsx_path: str, now_fn=lambda: datetime.now().isoformat(timespec="seconds")) -> ImportReport:
    with open(xlsx_path, "rb") as f:
        file_hash = hashlib.sha256(f.read()).hexdigest()
    try:
        wb = load_workbook(xlsx_path)
    except Exception:
        raise ApprovalStructureError("엑셀 파일을 열 수 없음")
    for name in (SHEET_SUMMARY, SHEET_CANDS, SHEET_META):
        if name not in wb.sheetnames:
            raise ApprovalStructureError(f"필수 시트 없음: {name}")
    ws, wc, wm = wb[SHEET_SUMMARY], wb[SHEET_CANDS], wb[SHEET_META]

    meta = {str(r[0].value): r[1].value for r in wm.iter_rows(min_col=1, max_col=2) if r[0].value is not None}
    meta_run = str(meta.get("run_id") or "")
    r_id = _find_label(ws, LBL_RUN_ID)
    r_app = _find_label(ws, LBL_RUN_APPROVAL)
    r_memo = _find_label(ws, LBL_RUN_MEMO)
    if not meta_run or r_id is None or r_app is None or r_memo is None:
        raise ApprovalStructureError("필수 칸 없음: run_id/실행 승인/승인 메모")
    if str(ws.cell(r_id, 2).value or "") != meta_run:
        raise ApprovalStructureError("run_id 불일치: 요약 시트와 메타 시트")
    if conn.execute("SELECT 1 FROM runs WHERE run_id=?", (meta_run,)).fetchone() is None:
        raise ApprovalStructureError("run_id 불일치: DB에 없는 실행")

    headers = {str(c.value).strip(): c.column for c in wc[1] if c.value is not None}
    for h in ("키워드", "도매처", ITEM_HEADER):
        if h not in headers:
            raise ApprovalStructureError(f"필수 칸 없음: 후보 시트 '{h}'")

    report = ImportReport(run_id=meta_run)
    valid_keys = {r["row_key"] for r in conn.execute("SELECT row_key FROM results WHERE run_id=?", (meta_run,))}
    events = []

    val = ws.cell(r_app, 2).value
    if val is not None and str(val).strip() != "":
        d = str(val).strip()
        if d in RUN_DECISIONS:
            memo = ws.cell(r_memo, 2).value
            events.append(("run", d, str(memo).strip() if memo is not None else ""))
        else:
            report.row_errors.append(f"요약 시트 실행 승인 값 오류: '{d}'")

    for row in range(2, wc.max_row + 1):
        mark = wc.cell(row, headers[ITEM_HEADER]).value
        if mark is None or str(mark).strip() == "":
            continue
        key = f"{wc.cell(row, headers['키워드']).value}|{wc.cell(row, headers['도매처']).value}"
        m = str(mark).strip().upper()
        if m not in ITEM_MAP:
            report.row_errors.append(f"후보 시트 {row}행 O/X 값 오류: '{mark}'")
        elif key not in valid_keys:
            report.row_errors.append(f"후보 시트 {row}행 알 수 없는 항목: {key}")
        else:
            events.append((f"item:{key}", ITEM_MAP[m], ""))

    ts = now_fn()
    for scope, decision, memo in events:
        cur = conn.execute(
            "INSERT OR IGNORE INTO approval_events (run_id, scope, decision, memo, event_at, source_hash) "
            "VALUES (?,?,?,?,?,?)", (meta_run, scope, decision, memo, ts, file_hash))
        if cur.rowcount:
            report.inserted += 1
        else:
            report.duplicates += 1
    conn.commit()
    return report
