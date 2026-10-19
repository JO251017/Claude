"""관리 엑셀(관리.xlsx) 하나로 입력·결과·승인을 관리한다.
입력 시트는 한글 머리글을 내부 키로 바꿔 읽고, 결과 시트는 실행할 때마다 새로 쓴다."""
import os
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment

from src.common.csvio import read_csv
from src.xlsx_style import BOX, INPUT, MUTED, NAVY, _fill, _font, _header

GUIDE = "안내"
SEEDS, KEYWORDS, COMPETITORS, COSTS, TRENDS = "씨앗", "키워드", "경쟁시세", "도매원가", "클릭추이"
DISCOVERY = "발굴 결과"
USE_HEADER = "키워드로 사용(O)"

# 입력 시트: (내부 키, 엑셀 머리글, 설명)
INPUT_SHEETS = {
    SEEDS: [("seed", "씨앗 키워드", "관심 분야의 넓은 단어. 0단계 발굴에 씀"),
            ("category", "분야 코드", "네이버쇼핑 카테고리 주소의 cat_id 숫자(데이터랩 자동 조회용)")],
    KEYWORDS: [("keyword", "키워드", "직접 검토할 키워드. 발굴 결과에서 O 표시한 키워드도 함께 검토됨"),
               ("category", "분야 코드", "네이버쇼핑 cat_id(데이터랩 자동 조회용)")],
    COMPETITORS: [("keyword", "키워드", ""), ("seller", "판매처", ""), ("price", "상품가", "원"),
                  ("shipping", "배송비", "고객이 내는 배송비. 무료배송은 0"),
                  ("researched_at", "조사일", "YYYY-MM-DD. 7일 넘으면 최신성 초과")],
    COSTS: [("keyword", "키워드", ""), ("product_name", "상품명", ""), ("cost", "도매 원가", "원"),
            ("shipping", "도매 배송비", "원"), ("supplier", "도매처", ""),
            ("ship_fee_rule", "배송비 규칙", "고정 / 수량당"),
            ("wholesale_date", "원가 확인일", "YYYY-MM-DD. 14일 넘으면 최신성 초과"),
            ("spec_info", "스펙 제공", "Y/N: 소재·규격을 글로 받을 수 있는지(AI 노출 준비)"),
            ("model_name", "모델명·브랜드", "AI 노출 준비"), ("dispatch_days", "출고 일수", "AI 노출 준비"),
            ("differentiator", "차별점", "한 줄(AI 노출 준비)")],
    TRENDS: [("keyword", "키워드", ""), ("month", "월", "YYYY-MM. 이번 달 제외 최근 15개월"),
             ("ratio", "데이터랩 값", "쇼핑인사이트 값. 데이터랩 자동 조회를 쓰면 비워 둠")],
}
CSV_NAMES = {SEEDS: "seeds.csv", KEYWORDS: "keywords.csv", COMPETITORS: "competitors.csv",
             COSTS: "costs.csv", TRENDS: "trends.csv"}

GUIDE_LINES = [
    ("스마트스토어 상품 관리 엑셀", ""),
    ("", ""),
    ("이 파일 하나로", "입력(노란 시트) → 실행 → 결과(파란 시트) 확인 → 승인까지 합니다. 실행할 때는 이 파일을 닫아 주세요."),
    ("", ""),
    ("0단계 키워드 발굴", "① '씨앗' 시트에 단어 입력  ② 0_키워드발굴 실행  ③ '발굴 결과'에서 쓸 키워드에 O"),
    ("1단계 순이익 후보표", "① '키워드'·'경쟁시세'·'도매원가' 입력(클릭추이는 직접 입력 또는 데이터랩 자동)  ② 1_후보표만들기 실행"),
    ("2단계 승인", "① '요약' 시트 맨 아래 '실행 승인' 선택, '후보' 시트에 O/X  ② 저장 후 2_승인가져오기 실행"),
    ("", ""),
    ("시트 색", "노란 탭 = 입력하는 곳, 파란 탭 = 프로그램이 채우는 곳(직접 고쳐도 다음 실행 때 새로 써짐)"),
    ("기준·요율", "config.yaml에서 바꿉니다(마진 15%, 최소 이익 2,000원 등). 순이익은 세전입니다."),
    ("기록", "실행마다 output 폴더에 결과 사본·로그가 남습니다."),
]
INPUT_TAB, OUTPUT_TAB = "F2C94C", "1F3A5F"


def _cell_text(key: str, v) -> str:
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m") if key == "month" else v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def read_sheet(wb, name: str) -> list[dict]:
    """한글 머리글 -> 내부 키. 시트가 없으면 빈 목록."""
    if name not in wb.sheetnames:
        return []
    spec = INPUT_SHEETS[name]
    ws = wb[name]
    header = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    col = {key: header.index(label) for key, label, _ in spec if label in header}
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        rec = {key: _cell_text(key, row[i]) if i < len(row) else "" for key, i in col.items()}
        if any(rec.values()):
            out.append(rec)
    return out


def selected_discovered(wb) -> list[dict]:
    """'발굴 결과'에서 O 표시한 키워드."""
    if DISCOVERY not in wb.sheetnames:
        return []
    ws = wb[DISCOVERY]
    header = [c.value for c in ws[1]]
    if USE_HEADER not in header or "키워드" not in header:
        return []
    iu, ik = header.index(USE_HEADER), header.index("키워드")
    ic = header.index("분야 코드") if "분야 코드" in header else None
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[iu] is not None and str(row[iu]).strip().upper() == "O" and row[ik]:
            out.append({"keyword": str(row[ik]).strip(), "category": _cell_text("category", row[ic]) if ic is not None else ""})
    return out


def keywords_for_run(wb) -> list[dict]:
    seen, out = set(), []
    for r in read_sheet(wb, KEYWORDS) + selected_discovered(wb):
        if r.get("keyword") and r["keyword"] not in seen:
            seen.add(r["keyword"])
            out.append(r)
    return out


def _input_sheet(wb, name: str, rows: list[dict]):
    ws = wb.create_sheet(name)
    ws.sheet_properties.tabColor = INPUT_TAB
    spec = INPUT_SHEETS[name]
    ws.append([label for _, label, _ in spec])
    for r in rows:
        ws.append([r.get(key, "") for key, _, _ in spec])
    _header(ws, 30)
    for i, (key, label, desc) in enumerate(spec, 1):
        c = ws.cell(1, i)
        if desc:
            c.comment = Comment(desc, "안내")
        ws.column_dimensions[c.column_letter].width = max(12, min(28, len(label) * 2 + 6))
    ws.freeze_panes = "A2"
    ws.page_setup.orientation = "landscape" if len(spec) > 5 else "portrait"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    for row in ws.iter_rows(min_row=2, max_row=max(ws.max_row, 30), max_col=len(spec)):
        for c in row:
            c.fill, c.border, c.font = _fill(INPUT), BOX, _font()
    return ws


def create_workbook(path: str, sample_dir: str | None = None) -> None:
    wb = Workbook()
    g = wb.active
    g.title = GUIDE
    g.sheet_view.showGridLines = False
    g.column_dimensions["A"].width = 22
    g.column_dimensions["B"].width = 100
    for a, b in GUIDE_LINES:
        g.append([a or None, b or None])          # 빈 문자열은 제목 글자가 옆 칸으로 넘치는 것을 막으므로 None
    g["A1"].font = _font(size=16, bold=True, color=NAVY)
    g.row_dimensions[1].height = 30
    g.page_setup.fitToWidth, g.page_setup.fitToHeight = 1, 0
    g.sheet_properties.pageSetUpPr.fitToPage = True
    for row in g.iter_rows(min_row=2):
        g.row_dimensions[row[0].row].height = 30 if row[1].value else 12
        row[0].alignment = Alignment(vertical="center")
        row[0].font = _font(bold=True, color=NAVY)
        row[1].font = _font(color="222B36" if row[0].value else MUTED)
        row[1].alignment = Alignment(wrap_text=True, vertical="center")
    for name in INPUT_SHEETS:
        rows = []
        if sample_dir and os.path.exists(os.path.join(sample_dir, CSV_NAMES[name])):
            rows = read_csv(os.path.join(sample_dir, CSV_NAMES[name]))
        _input_sheet(wb, name, rows)
    wb.save(path)


def open_workbook(path: str):
    return load_workbook(path)


def replace_sheet(wb, name: str, index: int | None = None):
    """같은 이름의 결과 시트를 지우고 새로 만든다(파란 탭)."""
    if name in wb.sheetnames:
        index = wb.sheetnames.index(name) if index is None else index
        del wb[name]
    ws = wb.create_sheet(name, index)
    ws.sheet_properties.tabColor = OUTPUT_TAB
    return ws


class WorkbookLockedError(Exception):
    pass


def save_workbook(wb, path: str) -> None:
    try:
        wb.save(path)
    except PermissionError:
        raise WorkbookLockedError(f"{os.path.basename(path)}이(가) 열려 있어 저장할 수 없습니다. 엑셀을 닫고 다시 실행하세요.") from None
